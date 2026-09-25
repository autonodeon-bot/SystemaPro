"""Appendices А–Д follow reference 26-2335; only source report facts are inserted."""
from io import BytesIO
from docx.shared import Cm, Inches, Pt
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from epb_layout import table, format_table, signature
from vessel_control_plan import method_rows, vessel_shells, shell_points, number, planned_count


def append_inspection_schemes(doc, ctx, add_picture):
    """Use the frozen report geometry when no usable uploaded scheme exists."""
    from copy import deepcopy
    from scheme_ndt_overlays import extract_geometry, render_all_layer_pngs

    g = ctx.g
    for key in ('control_scheme_image', 'control_scheme'):
        image = g(key, default=None)
        if isinstance(image, str) and not image.lower().endswith('.pdf'):
            if add_picture('', image, width_inches=8.6):
                return
    data = deepcopy(ctx.inspection_data.get('data') or {})
    shells = vessel_shells(g)
    if shells:
        elements = data.get('vessel_elements') or []
        data['vessel_elements'] = shells + [e for e in elements if isinstance(e, dict)
            and not any(s in str(e.get('name') or e.get('element_name') or '').lower()
                        for s in ('обечай', 'корпус'))]
    if not extract_geometry(data):
        doc.add_paragraph('Схема контроля не сформирована: отсутствуют геометрия и размеры сосуда в исходном отчёте.')
        return
    welds = data.get('weld_inspections') or []
    layers = {'vik', 'uzt'}
    for code, layer, rows in [('UZK', 'uzk', method_rows(welds, 'UZK')),
                              ('MPK', 'mpk', method_rows(welds, 'MPK')),
                              ('PVK', 'mpk', method_rows(welds, 'PVK')),
                              ('HARDNESS', 'hardness', data.get('hardness_tests'))]:
        if rows or planned_count(g, code):
            layers.add(layer)
    # Same canvas as the stored control-point coordinates used by the report.
    items = [item for item in render_all_layer_pngs(data, width=1800, height=1100)
             if item['layer'] in layers]
    if {item['layer'] for item in items} != layers:
        raise ValueError('Не удалось сформировать схемы контроля для ЭПБ')
    for index, item in enumerate(items):
        if index:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(BytesIO(item['png']), width=Inches(8.6))


def append_vessel_epb(doc, ctx, add_picture):
    g=ctx.g
    def heading(text):
        doc.add_page_break();doc.add_heading(text,level=1)
    def header(code):
        methods=[m for m in ctx.ndt_methods if str(m.get('method_code') or '').upper()==code]
        instrument=g(f'{code.lower()}_instruments',default=None)
        if not instrument:
            instrument='; '.join(str(m.get('equipment_name') or m.get('equipment') or '') for m in methods).strip('; ') or '—'
        rows=[('Заказчик',ctx.org),('Объект контроля',f'{ctx.device_name}, зав. № {ctx.serial}, рег. № {ctx.reg_no}'),
              ('Место проведения контроля',ctx.location),('Дата проведения',ctx.date_perf_ru),
              ('НТД, по которой выполнен контроль',g(f'{code.lower()}_normative', 'ndt_normative',default='—')),
              ('Приборы, инструменты и принадлежности',instrument),('Схема контроля','Приложение В')]
        t=doc.add_table(rows=0,cols=2);t.style='Table Grid'
        for values in rows:
            for c,text in zip(t.add_row().cells,values):c.text=str(text)
        format_table(t,widths=[4.3,13.4],header_rows=0,size=11)
        doc.add_paragraph()
    def act(no,title,code,first=False):
        if not first:doc.add_page_break()
        doc.add_heading(f'Акт № {no}. {title}',level=2);header(code)
    def conclusion(key):
        doc.add_paragraph('Заключение: '+str(g(key,default='Заключение не внесено.')))
        signature(doc,ctx)

    heading('Приложение А Акт о проведении работ по техническому диагностированию')
    doc.add_paragraph(f'{ctx.date_perf_ru} выполнено техническое диагностирование {ctx.device_name}, зав. № {ctx.serial}, рег. № {ctx.reg_no}.')
    doc.add_paragraph(f'Заказчик: {ctx.org}. Место эксплуатации: {ctx.location}.')
    doc.add_paragraph(f'Основание: {g("contract_number", "basis", default="—")}. Исходный технический отчёт: {g("source_report_number", default="—")}.')
    doc.add_paragraph('Состав и результаты работ:')
    for key,label in [('doc_analysis_conclusion','Анализ документации'),('vik_conclusion','ВИК'),('operational_conclusion','Оперативное диагностирование'),('uzt_conclusion','УЗТ'),('uzk_conclusion','УЗК'),('mpk_conclusion','МК'),('pvk_conclusion','ПВК'),('hardness_conclusion','Твердометрия'),('calculation_result','Расчётные процедуры')]:
        doc.add_paragraph(f'{label}: {g(key,default="сведения не внесены")}')
    signature(doc,ctx)

    heading('Приложение Б Акты по результатам неразрушающего контроля')
    act(1,'Визуального и измерительного контроля','VIK',first=True)
    doc.add_paragraph('Результаты наружного и внутреннего осмотра')
    doc.add_paragraph(str(g('vik_conclusion',default='Результаты осмотра не внесены.')))
    defects=g('visual_defects',default=[]) or []
    if defects:table(doc,['Элемент','Дефект','Размеры'],[[r.get('location'),r.get('description'),r.get('size')] for r in defects],[5,9,3.7])
    oval=g('ovality_measurements',default=[]) or []
    if oval:
        doc.add_paragraph('Результаты измерительного контроля. Определение овальности')
        doc.add_paragraph('a = 200 × (Dmax − Dmin) / (Dmax + Dmin), %.')
        table(doc,['Сечение','Dmin, мм','Dmax, мм','Фактическая овальность, %','Допустимая овальность, %'],
              [[r.get('section_number'),r.get('min_diameter',r.get('diameter_min')),r.get('max_diameter',r.get('diameter_max')),r.get('deviation_percent',r.get('ovality_percent')),r.get('allowed_ovality') or g('allowed_ovality',default='—')] for r in oval],[2,3,3,4.8,4.9])
    deflections=g('deflection_measurements',default=[]) or []
    if deflections:
        doc.add_paragraph('Измерение отклонений от прямолинейности образующей')
        table(doc,['Сечение','Прогиб, мм','Прогиб, %'],[[r.get('section_number'),r.get('deflection_mm'),r.get('deflection_percent')] for r in deflections],[5,6,6.7])
    signature(doc,ctx)
    act(2,'Оперативного (функционального) диагностирования','FUNCTIONAL')
    rows=g('functional_checks',default=[]) or []
    table(doc,['Контролируемый параметр','Результат'],[[r.get('name'),r.get('result') or r.get('description')] for r in rows] or [['Сведения не предоставлены','—']],[7,10.7])
    for key,label in [('functional_params','Рабочие параметры'),('vibration','Вибрация'),('support_state','Опоры'),('kip_state','КИП'),('manometer_verification','Поверка манометра')]:
        value=g(key,default=None)
        if value:doc.add_paragraph(f'{label}: {value}')
    conclusion('operational_conclusion')

    act(3,'Ультразвукового контроля толщины стенок элементов сосуда','UZT')
    doc.add_paragraph('Результаты контроля')
    grouped,other=shell_points(vessel_shells(g),g('thickness_measurements',default=[]) or [])
    for name,points in grouped.items():
        present={number(p.get('point_number') or p.get('section_number')) for p in points}
        points.extend({'point_number':i,'thickness':None} for i in range(1,13) if number(i) not in present)
        points.sort(key=lambda p:number(p.get('point_number') or p.get('section_number')) or 0)
    grouped.update(other)
    t=doc.add_table(rows=2,cols=12);t.style='Table Grid';t.autofit=False
    t.cell(0,0).merge(t.cell(1,0)).text='Наименование элемента сосуда'
    t.cell(0,1).merge(t.cell(1,1)).text='Номинальная толщина, мм'
    t.cell(0,2).merge(t.cell(0,11)).text='Результаты измерений'
    for i in range(5):t.cell(1,2+i*2).text='№ точки';t.cell(1,3+i*2).text='Толщина, мм'
    for name,points in grouped.items():
        start=len(t.rows)
        nominal=next((p.get('nominal_thickness') for p in points if p.get('nominal_thickness') not in (None,'')),g('wall_thickness',default='—'))
        for offset in range(0,len(points),5):
            cells=t.add_row().cells
            for j,p in enumerate(points[offset:offset+5]):
                cells[2+2*j].text=str(p.get('point_number') or p.get('section_number') or offset+j+1)
                val=p.get('thickness');cells[3+2*j].text=str(val).replace('.',',') if val is not None else '—'
        end=len(t.rows)-1
        if end>=start:
            t.cell(start,0).merge(t.cell(end,0)).text=name
            t.cell(start,1).merge(t.cell(end,1)).text=str(nominal).replace('.',',')
    format_table(t,widths=[3,1.7]+[1.05,1.55]*5,header_rows=2,size=9)
    for name,points in grouped.items():
        values=[number(p.get('thickness')) for p in points if number(p.get('thickness')) is not None]
        if values:
            minimum=min(values)
            for row in t.rows[2:]:
                if row.cells[0].text==name:
                    for j in range(3,12,2):
                        if number(row.cells[j].text)==minimum:
                            for run in row.cells[j].paragraphs[0].runs:run.bold=True
    conclusion('uzt_conclusion')

    welds=g('weld_inspections',default=[]) or []
    surface=[('MPK','Магнитопорошкового контроля','mpk_conclusion'),('PVK','Капиллярного контроля (ПВК)','pvk_conclusion')]
    selected=[x for x in surface if method_rows(welds,x[0]) or planned_count(g,x[0]) not in (None,0)]
    for index,(code,title,key) in enumerate(selected or surface[:1]):
        act(4 if index==0 else 7,title,code)
        rows=method_rows(welds,code)
        if rows:table(doc,['Номер соединения / участка','Выявленный дефект','Описание дефекта','Заключение'],[[r.get('weld_number') or r.get('location'),r.get('defect_number') or r.get('defect_description'),r.get('description'),r.get('conclusion')] for r in rows],[4.1,4,6.1,3.5])
        else:doc.add_paragraph('Результаты контроля не предоставлены.')
        conclusion(key)
    act(5,'Ультразвукового контроля сварных соединений','UZK')
    rows=method_rows(welds,'UZK')
    if rows:
        table(doc,['Номер соединения','Номер дефекта','Эквивалентная площадь, мм²','Глубина, мм','Протяжённость, мм','Форма дефекта','Место, мм','Заключение'],
            [[r.get('weld_number'),r.get('defect_number'),r.get('equivalent_area'),r.get('depth'),r.get('length'),r.get('defect_type') or r.get('defect_description'),r.get('defect_location'),r.get('conclusion')] for r in rows],[2.5,1.7,2.5,1.7,2.3,3.3,1.7,2])
    else:doc.add_paragraph('Результаты контроля не предоставлены.')
    conclusion('uzk_conclusion')
    act(6,'Оценка механических свойств элементов сосуда','HARDNESS')
    hardness=g('hardness_tests',default=[]) or []
    if hardness:
        ht=doc.add_table(rows=3,cols=8);ht.style='Table Grid'
        for c,title in enumerate(['Соединение / участок','Допустимая твёрдость основного металла, HB','Шов и ЗТВ, HB, не более']):
            ht.cell(0,c).merge(ht.cell(2,c)).text=title
        ht.cell(0,3).merge(ht.cell(0,7)).text='Фактическая твёрдость, HB'
        ht.cell(1,3).merge(ht.cell(1,4)).text='Основной металл'
        ht.cell(1,5).text='Шов';ht.cell(1,6).merge(ht.cell(1,7)).text='Зона термического влияния'
        for c,label in enumerate(['Т1','Т5','Т3','Т2','Т4'],3):ht.cell(2,c).text=label
        for r in hardness:
            values=[r.get('weld_number') or r.get('location'),r.get('allowed_hardness_base'),r.get('allowed_hardness_weld'),r.get('hardness_base_t1') or r.get('hardness_base'),r.get('hardness_base_t5'),r.get('hardness_weld'),r.get('hardness_haz_t2') or r.get('hardness_haz'),r.get('hardness_haz_t4')]
            for c,value in zip(ht.add_row().cells,values):c.text=str(value if value is not None else '—')
        format_table(ht,[3,3.1,3.1,1.7,1.7,1.7,1.7,1.7],header_rows=3)
    else:doc.add_paragraph('Результаты твердометрии не предоставлены.')
    conclusion('hardness_conclusion')

    sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width=Cm(29.7);sec.page_height=Cm(21);sec.left_margin=Cm(1.27);sec.right_margin=Cm(.63)
    doc.add_heading('Приложение В Схема проведения неразрушающего контроля',level=1)
    append_inspection_schemes(doc, ctx, add_picture)
    doc.add_paragraph('На каждой обечайке — 12 точек УЗТ. Обозначения элементов и точек соответствуют акту № 3. Участки УЗК, МК (ПВК) и твердометрии — по техническому заданию.')
    sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.PORTRAIT
    sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=Cm(2.25);sec.right_margin=Cm(1)
    doc.add_heading('Приложение Г Расчетные и аналитические процедуры оценки и прогнозирования технического состояния объекта экспертизы',level=1)
    calc=g('calculation_data',default={}) or {}
    doc.add_paragraph('Расчёт на прочность и определение остаточного ресурса сосуда')
    doc.add_paragraph(str(calc.get('method') or 'Методика расчёта не указана.'))
    elements=calc.get('elements') or []
    if elements:
        doc.add_paragraph('Таблица 1. Исходные данные и результаты расчёта на прочность')
        table(doc,['Элемент','Номинальная толщина, мм','Минимальная измеренная толщина, мм','Минимально допустимая толщина, мм','Допустимое давление, МПа','Заключение'],
            [[r.get('element'),r.get('nominal_thickness'),r.get('min_thickness'),r.get('min_allowed_thickness'),r.get('allowed_pressure'),r.get('conclusion')] for r in elements],[3.2,2.5,3,3,2.4,3.6])
        doc.add_paragraph('Таблица 2. Расчёт остаточного ресурса сосуда по элементам')
        table(doc,['Элемент','Скорость коррозии, мм/год','Остаточный ресурс, лет'],[[r.get('element'),r.get('corrosion_rate'),r.get('residual_life_years')] for r in elements],[7.7,5,5])
    else:
        labels={'residual_life_years':'Остаточный ресурс, лет','min_allowed_thickness':'Минимальная толщина, мм','allowed_pressure':'Допустимое давление, МПа','corrosion_rate':'Скорость коррозии, мм/год','description':'Описание','conclusion':'Заключение','method':'Методика','min_thickness':'Минимальная измеренная толщина, мм'}
        table(doc,['Расчётная величина','Значение'],[[labels.get(k,k),v] for k,v in calc.items() if not isinstance(v,(dict,list))] or [['Данные не предоставлены','—']],[9,8.7])
    doc.add_paragraph(str(g('calculation_result',default='Заключение по расчётным процедурам не внесено.')))
    signature(doc,ctx)
    heading('Приложение Д Акт проведения гидравлических испытаний')
    table(doc,['Наименование','Сведения'],[['Техническое устройство',f'{ctx.device_name}, зав. № {ctx.serial}, рег. № {ctx.reg_no}'],['Организация',ctx.org],['Место эксплуатации',ctx.location]],[5,12.7])
    hydro=g('hydrostatic_test',default={}) or {}
    if not hydro:
        history=g('hydraulic_test_history',default=[]) or [];hydro=history[-1] if history else {}
    for key,label in [('date','Дата'),('test_pressure','Пробное давление, МПа'),('medium','Среда'),('hold_time','Продолжительность выдержки'),('result','Результат')]:
        doc.add_paragraph(f'{label}: {hydro.get(key) or (hydro.get("pressure") if key=="test_pressure" else None) or "—"}')
    doc.add_paragraph('Представители эксплуатирующей организации:')
    doc.add_paragraph('Ответственный за производственный контроль __________________')
    doc.add_paragraph('Ответственный за исправное состояние и безопасную эксплуатацию __________________')
