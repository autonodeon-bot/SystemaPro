"""Dynamic EPB layout derived from the retained 26-2335 vessel reference."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def new_epb_document():
    doc = Document(Path(__file__).parent / 'report_forms' / 'epb_vessel_layout.docx')
    names = {s.name for s in doc.styles}
    for name in ['Normal', 'Heading 1', 'Heading 2', 'Heading 3']:
        # Some source styles use localized identifiers inaccessible by English lookup.
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = 'Arial'; style.font.size = Pt(12); style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.line_spacing = 1
        if name.startswith('Heading'):
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(10)
    try:
        doc.styles['Table Grid']
    except KeyError:
        style = doc.styles.add_style('Table Grid', WD_STYLE_TYPE.TABLE)
        pr = OxmlElement('w:tblPr'); borders = OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            node=OxmlElement('w:'+side);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');borders.append(node)
        pr.append(borders);style.element.append(pr)
    sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.left_margin=Cm(2.25);sec.right_margin=Cm(1);sec.top_margin=Cm(.63);sec.bottom_margin=Cm(2)
    sec.different_first_page_header_footer=True
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
    return doc


def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers));t.style='Table Grid';t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,text in zip(t.rows[0].cells,headers):c.text=str(text)
    for values in rows:
        for c,value in zip(t.add_row().cells,values):c.text=str(value if value not in (None,'') else '—')
    format_table(t, widths=widths)
    return t


def format_table(t, widths=None, header_rows=1, size=10):
    if widths:
        t.autofit=False
        for column,width in zip(t.columns,widths):column.width=Cm(width)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');pr.append(no_split)
        if i<header_rows:
            repeat=OxmlElement('w:tblHeader');pr.append(repeat)
        seen=set()
        for j,c in enumerate(row.cells):
            if c._tc in seen:continue
            seen.add(c._tc)
            if widths and j<len(widths):
                span=c._tc.grid_span
                c.width=Cm(sum(widths[j:j+span]))
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
                p.paragraph_format.line_spacing=1
                p.paragraph_format.keep_with_next=i<header_rows
                for r in p.runs:r.font.name='Arial';r.font.size=Pt(size);r.bold=i<header_rows


def signature(doc, ctx):
    doc.add_paragraph('Контроль выполнил и заключение по результатам НК выдал:')
    names=ctx.inspectors or [e.get('full_name') for e in ctx.inspection_engineers if e.get('full_name')] or ['__________________']
    p=doc.add_paragraph(f'Специалист НК II уровня    _________________    {names[0]}')
    p.paragraph_format.keep_with_next=True
    doc.add_paragraph(f'{ctx.director_title}    _________________    {ctx.director_name}')


def finalize_epb(doc):
    """Populate real TOC page references without copying the reference page numbers."""
    headings=[p for p in doc.paragraphs if p.style.name=='Heading 1' and p.text!='Содержание']
    for i,p in enumerate(headings):
        name=f'epb_section_{i}'
        start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(500+i));start.set(qn('w:name'),name)
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(500+i))
        p._p.insert(0,start);p._p.append(end)
        for candidate in doc.paragraphs:
            if candidate is p or candidate._p is p._p or candidate.style.name=='Heading 1':continue
            if candidate.text==p.text:
                from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
                candidate.paragraph_format.tab_stops.add_tab_stop(Cm(17.4),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
                candidate.add_run('\t')
                field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),f'PAGEREF {name} \\h')
                candidate._p.append(field)
                break
    for p in doc.paragraphs:
        if p.text.startswith('Таблица '):
            p.paragraph_format.keep_with_next=True
