"""User-entered supplemental values, shared by every official equipment form."""
from typing import Any


def append_report_fields(doc: Any, data: dict) -> int:
    additional = data.get("additional_data")
    if not isinstance(additional, dict):
        return 0
    rows = additional.get("report_fields")
    rows = list(rows) if isinstance(rows, list) else []
    from report_schema import schema_for_form
    profile_values = additional.get("profile_values")
    if isinstance(profile_values, dict):
        for field in schema_for_form(data.get("report_form_id") or "to-1")["fields"]:
            value = profile_values.get(field["key"])
            if value is not None and str(value).strip():
                rows.append({"section": "Характеристики оборудования", "label": field["label"], "value": value, "unit": field["unit"]})
    values = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        cells = [str(row.get(key) if row.get(key) is not None else "").strip()
                 for key in ("section", "label", "value", "unit")]
        if not any(cells):
            continue
        values.append(cells)
    if not values:
        return append_profile_measurements(doc, data)
    doc.add_page_break()
    doc.add_heading("Дополнительные сведения обследования", level=1)
    doc.add_paragraph("Сведения, внесённые при обследовании дополнительно к основной форме отчёта.")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ("Раздел / протокол", "Показатель", "Значение / результат", "Ед. изм.")):
        cell.text = label
    for cells in values:
        for cell, value in zip(table.add_row().cells, cells):
            cell.text = value or "Не указано"
    return len(values) + append_profile_measurements(doc, data)


def append_profile_measurements(doc, data):
    from report_schema import schema_for_form
    rows = (data.get('additional_data') or {}).get('profile_measurements') or []
    columns = schema_for_form(data.get('report_form_id') or 'to-1')['measurement_columns']
    rows = [row for row in rows if isinstance(row, dict) and any(row.get(c['key']) not in (None, '') for c in columns)]
    if not rows or not columns:
        return 0
    doc.add_page_break()
    doc.add_heading('Профильные измерения', level=1)
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = 'Table Grid'
    from docx.oxml import OxmlElement
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for cell, column in zip(table.rows[0].cells, columns):
        cell.text = column['label'] + (f", {column['unit']}" if column['unit'] else '')
    for row in rows:
        for cell, column in zip(table.add_row().cells, columns):
            value = row.get(column['key'])
            cell.text = str(value) if value not in (None, '') else 'Не указано'
    return len(rows)
