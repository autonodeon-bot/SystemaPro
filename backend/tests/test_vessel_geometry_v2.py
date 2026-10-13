import copy
import json
from pathlib import Path
import pytest
from vessel_geometry_v2 import normalize_v2, position
from vessel_scheme_renderer import normalize_geometry, render_vessel_scheme

@pytest.fixture
def geometry():
    return json.loads((Path(__file__).parents[2]/'artifacts/vessel-constructor-v2/fixture.json').read_text(encoding='utf-8'))

def test_shared_fixture_roundtrip_and_coordinates(geometry):
    original=copy.deepcopy(geometry)
    g=normalize_geometry(geometry)
    assert geometry==original
    assert len(g['welds'])==7
    assert len(g['control_points'])==78
    assert g['control_zones'][0]['coverage']['span_deg']==20
    e=g['elements'][0]
    owner=next(p for p in g['parts'] if p['id']==e['parent_id'])
    assert position(e,owner,90)==(.175,0)
    g['future_extension']={'foo':1}
    assert normalize_geometry(g)['future_extension']=={'foo':1}

@pytest.mark.parametrize('orientation',['vertical','horizontal'])
def test_render_does_not_generate_measurements(geometry,orientation):
    geometry['orientation']=orientation
    png,g,points=render_vessel_scheme(geometry,width=1190,height=842)
    assert png.startswith(b'\x89PNG') and not points
    assert g['elements'][0]['placement']==geometry['elements'][0]['placement']
    (Path(__file__).parents[2]/f'artifacts/vessel-constructor-v2/server-{orientation}.png').write_bytes(png)

@pytest.mark.parametrize('bad',['nan','owner','duplicate','zone','negative'])
def test_rejects_invalid_model(geometry,bad):
    if bad=='nan':geometry['elements'][0]['placement']['theta_deg']=float('nan')
    if bad=='owner':geometry['elements'][0]['parent_id']='missing'
    if bad=='duplicate':geometry['parts'][1]['id']=geometry['parts'][0]['id']
    if bad=='zone':geometry['control_zones'][0]['weld_id']='missing'
    if bad=='negative':geometry['parts'][2]['dimensions']['length_mm']=-1
    with pytest.raises(ValueError):normalize_v2(geometry)

def test_cover_detail(geometry):
    from vessel_geometry_v2 import render_v2
    cover=next(p for p in geometry['parts'] if p['kind']=='cover')
    png,_,_=render_v2(normalize_v2(geometry),1190,842,detail=cover['id'])
    assert png.startswith(b'\x89PNG')


def test_vector_pdf_and_svg(geometry):
    import fitz
    from vessel_vector import export_pdf
    from vessel_geometry_v2 import render_v2
    doc = fitz.open(stream=export_pdf(geometry), filetype='pdf')
    assert len(doc) >= len(geometry['parts']) + 1
    assert not doc[0].get_images()
    assert len(doc[0].get_drawings()) > 30
    assert 'Сосуд' in doc[0].get_text()
    assert b'<svg' in render_v2(geometry,1190,842,svg=True)[0]


def test_bind_plan_keeps_measured_values(geometry):
    from vessel_snapshot_api import bind_plan
    point = geometry['control_points'][0]
    data = {'thickness_measurements': [{'scheme_point_id':point['id'], 'thickness':10.25}]}
    bound = bind_plan(data, geometry)
    assert bound['thickness_measurements'][0]['thickness'] == 10.25
    assert len(bound['thickness_measurements']) == len([p for p in geometry['control_points'] if not p.get('archived')])
    assert bind_plan(bound, geometry) == bound


def test_idempotence_ignores_generated_images(geometry):
    from vessel_snapshot_guard import guard_vessel_update
    geometry['operation_id'] = 'render-operation'
    old = {'base_vessel_scheme': {'geometry':copy.deepcopy(geometry)}}
    old['base_vessel_scheme']['geometry']['views'].append({'id':'generated', 'kind':'part_detail','image_path':'server.png'})
    assert guard_vessel_update(old, {'base_vessel_scheme':{'geometry':geometry}})


def test_legacy_layer_codes_are_mapped_for_v2(geometry):
    uzt=render_vessel_scheme(geometry,width=1190,height=842,scheme_layer='uzt')[0]
    vik=render_vessel_scheme(geometry,width=1190,height=842,scheme_layer='vik')[0]
    assert uzt != vik


def test_v2_report_keeps_details_and_independent_pvk(geometry):
    from docx import Document
    from form_template_filler import _insert_schemes_and_photos
    doc=Document()
    data={'base_vessel_scheme':{'geometry':geometry},'report_form_id':'to-1',
          'weld_inspections':[{'control_method':'PVK','weld_number':'К1','defects':'PVK-unique-result','result_status':'recorded'},
                              {'control_method':'MPK','weld_number':'К1','defects':'MPK-independent','result_status':'recorded'}]}
    _insert_schemes_and_photos(doc,{'data':data,'attachments':{},'find_image':lambda p:p})
    assert len(doc.inline_shapes)>=len(geometry['parts'])+5
    assert 'Крышка' in '\n'.join(p.text for p in doc.paragraphs)
    assert 'PVK-unique-result' in '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)

def test_revision_conflict_and_old_client_guard(geometry):
    from fastapi import HTTPException
    from vessel_snapshot_guard import guard_vessel_update
    old={'base_vessel_scheme':{'geometry':geometry}}
    new=copy.deepcopy(old)
    new['base_vessel_scheme']['geometry'].update(base_revision=geometry['revision'],revision=geometry['revision']+1,operation_id='op-1')
    saved=guard_vessel_update(old,new)
    assert guard_vessel_update(saved,new)==saved
    stale=copy.deepcopy(new);stale['base_vessel_scheme']['geometry']['operation_id']='op-2'
    with pytest.raises(HTTPException) as conflict:guard_vessel_update(saved,stale)
    assert conflict.value.status_code==409
    with pytest.raises(HTTPException):guard_vessel_update(saved,{'base_vessel_scheme':{'geometry':{'shell_count':3}}})

def test_v2_scope_uses_ids_not_global_numbers(geometry):
    from vessel_control_plan import control_scope
    points=geometry['control_points']
    data={'base_vessel_scheme':{'geometry':geometry},'thickness_measurements':[{'scheme_point_id':p['id'],'thickness':10,'point_number':p.get('body_number',p['local_number'])} for p in points]}
    def getter(*keys,default=None):
        return next((data[k] for k in keys if k in data),default)
    shells,scope,issues=control_scope(getter)
    assert len(shells)==6
    assert not any('точ' in issue for issue in issues)

def test_planned_weld_rows_are_not_completed_control(geometry):
    from vessel_control_plan import method_rows, planned_count
    rows=[{'scheme_zone_id':'z1','control_method':'MPK','result_status':'planned'},
          {'scheme_zone_id':'z2','control_method':'UZK','result_status':'recorded'}]
    assert method_rows(rows,'MPK')==[]
    assert len(method_rows(rows,'UZK'))==1
    assert method_rows(rows,'PVK')==[]
    data={'base_vessel_scheme':{'geometry':geometry}}
    assert planned_count(lambda *keys,default=None:next((data[k] for k in keys if k in data),default),'MPK')==1
