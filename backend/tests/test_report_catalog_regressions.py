from docx import Document
import pytest
from report_forms_registry import EQUIPMENT_TYPE_TO_FORM, suggest_form_id, resolve_form_path
from form_template_filler_generic import _fill_kv_tables


@pytest.mark.parametrize("code, expected", EQUIPMENT_TYPE_TO_FORM.items())
def test_explicit_equipment_type_selects_its_official_form(code, expected):
    assert suggest_form_id(code, "Объект") == expected


@pytest.mark.parametrize("code, name, expected", [
    ("CRANE_RUNWAY", "Подкрановые пути", "to-2"),
    ("", "Подкрановые пути", "to-2"),
    ("PIPELINE_VALVE", "Кран шаровой", "to-24"),
    ("PIPELINE", "Подземный трубопровод", "to-33"),
    ("", "Фонтанная арматура", "to-27"),
    ("", "Факельная установка", "to-44"),
])
def test_specific_forms_are_not_replaced_by_generic_name_matches(code, name, expected):
    assert suggest_form_id(code, name) == expected


def test_three_column_pressure_table_preserves_labels_and_zero_values():
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    for row, qualifier in zip(table.rows, ["рабочее", "расчётное", "пробное"]):
        row.cells[0].text = "Давление"
        row.cells[1].text = qualifier
    _fill_kv_tables(doc, {"working_pressure": 0, "design_pressure": "1,6", "test_pressure": "2,0"})
    assert [r.cells[2].text for r in table.rows] == ["0", "1,6", "2,0"]
    assert table.cell(1, 1).text == "расчётное"


def test_unknown_template_cannot_escape_catalog():
    assert resolve_form_path("../report_forms/to-1") is None


@pytest.mark.parametrize('form_id', ['to-5', 'to-8', 'to-12', 'to-24', 'to-44'])
def test_non_vessel_without_geometry_never_gets_default_vessel_schemes(form_id, monkeypatch):
    from unittest.mock import Mock
    from form_template_filler import insert_ndt_layer_schemes
    renderer = Mock(return_value=[])
    monkeypatch.setattr('form_template_filler.render_all_layer_pngs', renderer)
    assert insert_ndt_layer_schemes(Document(), {'report_form_id': form_id}, {}, None, kind=form_id) == 0
    renderer.assert_not_called()


def test_authored_non_vessel_scheme_receives_correct_kind(monkeypatch):
    from unittest.mock import Mock
    from form_template_filler import insert_ndt_layer_schemes
    renderer = Mock(return_value=[])
    monkeypatch.setattr('form_template_filler.render_all_layer_pngs', renderer)
    insert_ndt_layer_schemes(Document(), {'report_form_id': 'to-5', 'scheme_geometry': {'width': 200}}, {}, None)
    assert renderer.call_args.args[0]['equipment_kind'] == 'transformer'


def test_zero_thickness_is_not_silently_dropped():
    from form_template_filler_generic import _fill_uzt_tables
    doc = Document()
    table = doc.add_table(rows=2, cols=5)
    table.cell(0, 0).text = 'Точка замера толщины'
    _fill_uzt_tables(doc, {'thickness_measurements': [{'nominal': 0, 'thickness': 0}]})
    assert table.cell(1, 3).text == '0'
    assert table.cell(1, 4).text == '0'


def test_merged_cover_labels_are_preserved_and_values_are_filled():
    doc = Document()
    table = doc.add_table(rows=4, cols=4)
    for row, label in zip(table.rows, ['наименование', 'заводской №', 'ОПО регистрационный №', 'класс опасности ОПО']):
        row.cells[0].merge(row.cells[1]).text = label
        row.cells[2].merge(row.cells[3])
    _fill_kv_tables(doc, {'name': 'Компрессор', 'serial_number': '123', 'reg_number': 'WRONG', 'opo_reg_number': 'ОПО-456', 'opo_hazard_class': 'III'})
    assert [row.cells[2].text for row in table.rows] == ['Компрессор', '123', 'ОПО-456', 'III']
    assert table.cell(1, 0).text == 'заводской №'
