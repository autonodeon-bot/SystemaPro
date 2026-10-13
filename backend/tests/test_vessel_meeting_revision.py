import copy
import json
from pathlib import Path
import pymupdf
import pytest
from vessel_sheet import sheet_spec, main_element_rows
from vessel_vector import export_pdf
from vessel_geometry_v2 import position, render_v2

def example():
    return json.loads((Path(__file__).parents[2]/'public/examples/visio-vessel.json').read_text(encoding='utf-8'))

@pytest.mark.parametrize('orientation', ['horizontal','vertical'])
def test_adaptive_sheet_is_one_page_and_preserves_marks(orientation):
    g=example();g['orientation']=orientation
    spec=sheet_spec(g)
    assert spec['paper']=='A3'
    doc=pymupdf.open(stream=export_pdf(g),filetype='pdf')
    assert len(doc)==1
    page=doc[0]
    assert (page.rect.width>page.rect.height)==(orientation=='horizontal')
    assert abs(max(page.rect.width,page.rect.height)-420*72/25.4)<1
    svg=render_v2(g,spec['width'],spec['height'],svg=True)[0].decode()
    for point in g['control_points']:assert point['id'] in svg
    assert 'твердометрия' in page.get_text()
    g['print_settings']={'paper':'A4','detail_ids':[g['parts'][0]['id']]}
    assert len(pymupdf.open(stream=export_pdf(g),filetype='pdf'))==2

def test_grouping_does_not_merge_different_materials_or_include_nozzles():
    rows=[dict(name=f'Обечайка {i}',element_type='shell',material='09Г2С',nominal_thickness=12) for i in range(6)]
    rows += [dict(name='Обечайка 7',element_type='shell',material='20',nominal_thickness=10),dict(name='Правое днище',element_type='head'),dict(name='Э1',element_type='nozzle')]
    before=copy.deepcopy(rows);grouped=main_element_rows(rows,{'shell_length':12000})
    assert rows==before
    assert len(grouped)==3 and grouped[0]['quantity']==6
    assert grouped[0]['length_mm']==''  # no total passport length assigned to a subgroup

@pytest.mark.parametrize('side,angle,expected',[('A',0,(0,.5)),('A',90,(.5,0)),('B',0,(1,.5)),('B',180,(0,.5))])
def test_explicit_head_convention_and_legacy_preservation(side,angle,expected):
    owner={'kind':'head','side':side,'dimensions':{'diameter_mm':2000}}
    point={'placement':{'mode':'physical','radius_projection_mm':1000,'theta_deg':angle,'angle_convention':'visio-head-2026'}}
    assert position(point,owner)==pytest.approx(expected)
    point['placement'].pop('angle_convention');point['placement']['theta_deg']=0
    assert position(point,owner)==pytest.approx((.5,0))

def test_head_names_follow_orientation_without_changing_identity():
    from vessel_sheet import orient_head_names, orient_report_names
    g=example();before=[p['id'] for p in g['parts']]
    g['orientation']='vertical';orient_head_names(g)
    assert [p['designation'] for p in g['parts'] if p['kind']=='head']==['Верхнее днище','Нижнее днище']
    assert [p['id'] for p in g['parts']]==before
    data={'base_vessel_scheme':{'geometry':g},'thickness_measurements':[{'location':'Левое днище','thickness':12,'scheme_point_id':'stable'}]}
    updated=orient_report_names(data)
    assert updated['thickness_measurements'][0]=={'location':'Верхнее днище','thickness':12,'scheme_point_id':'stable'}
    assert data['thickness_measurements'][0]['location']=='Левое днище'
    g['orientation']='horizontal';orient_head_names(g)
    assert [p['designation'] for p in g['parts'] if p['kind']=='head']==['Левое днище','Правое днище']

def test_materials_table_contains_only_grouped_shell_and_two_heads():
    from docx import Document
    from form_template_filler import _fill_materials
    elements=[{'name':f'Обечайка {i}','material':'09Г2С','yield_strength':'345'} for i in range(1,7)]
    elements += [{'name':name,'material':'09Г2С'} for name in ['Левое днище','Правое днище','Крышка люка Т','Штуцер Н1']]
    data={'base_vessel_scheme':{'geometry':example()},'vessel_elements':elements}
    def get(*keys,default=None):return next((data[k] for k in keys if k in data),default)
    doc=Document();table=doc.add_table(rows=5,cols=10)
    _fill_materials(table,{'data':data,'g':get})
    assert [r.cells[0].text for r in table.rows[2:]]==['Обечайка','Левое днище','Правое днище']
    assert table.rows[2].cells[3].text=='345'
    # Explicit material rows use the same grouping; differing properties stay distinct.
    data['materials']=[{'element':'Обечайка 1','grade':'09Г2С','yield_strength':'345'},{'element':'Обечайка 2','grade':'09Г2С','yield_strength':'300'},{'element':'Штуцер Н1','grade':'09Г2С'}]
    doc=Document();table=doc.add_table(rows=4,cols=10);_fill_materials(table,{'data':data,'g':get})
    assert [r.cells[3].text for r in table.rows[2:]]==['345','300']
