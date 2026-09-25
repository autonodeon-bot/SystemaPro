import hashlib
import json

import pytest
from docx import Document
from epb_protocol_tables import append_epb_protocol_uzt, append_epb_protocol_weld_control
from report_snapshot import load_technical_report_snapshot
from vessel_control_plan import vessel_shells, control_scope, method_rows, planned_count


def getter(data):
    def g(*keys, default=None):
        return next((data[k] for k in keys if data.get(k) not in (None, '')), default)
    return g


@pytest.mark.parametrize('length,lengths', [(11000, [2000]*5+[1000]), (4000, [2000, 2000]), (900, [900])])
def test_shell_segmentation(length, lengths):
    assert [s['length_mm'] for s in vessel_shells(getter({'body_length_mm': length}))] == lengths


def test_passport_overrides_assumption():
    shells = [{'name': 'Обечайка 1', 'length_mm': 1500}, {'name': 'Обечайка 2', 'length_mm': 9500}]
    assert vessel_shells(getter({'body_length_mm': 11000, 'vessel_elements': shells})) == shells


def test_aggregate_passport_shell_is_split():
    shells = vessel_shells(getter({'vessel_elements': [{'name': 'Обечайка', 'length_mm': 11000, 'material': '09Г2С'}]}))
    assert len(shells) == 6
    assert shells[-1]['length_mm'] == 1000
    assert all(s['material'] == '09Г2С' for s in shells)


def test_missing_measurements_stay_missing_and_all_72_points_are_printed():
    data = {'body_length_mm': 11000}
    doc = Document()
    append_epb_protocol_uzt(doc, getter(data), lambda: None, [])
    assert len(doc.tables) == 6
    assert all(len(t.rows) == 4 for t in doc.tables)
    assert all(t.rows[-1].cells[-2].text == '12' for t in doc.tables)
    assert all(t.rows[-1].cells[-1].text == '—' for t in doc.tables)
    assert 'не выявлено' not in '\n'.join(p.text for p in doc.paragraphs)


def test_scope_checks_unique_points_and_method_tags():
    points = [{'location': 'Обечайка 1', 'point_number': i, 'thickness': '9,5'} for i in range(1, 13)]
    data = {'body_length_mm': 2000, 'thickness_measurements': points,
            'technical_task': {'uzk_sections': 1, 'mpk_sections': 0, 'pvk_sections': 1, 'hardness_sections': 0},
            'weld_inspections': [{'method': 'УЗК'}, {'method': 'ПВК'}]}
    assert not control_scope(getter(data))[2]
    data['thickness_measurements'] = points[:-1] + [points[0]]
    assert any('11 из 12' in s for s in control_scope(getter(data))[2])
    assert method_rows([{}, {'method': 'ПВК'}], 'MPK') == []


@pytest.mark.parametrize('bad', [-1, 1.5, 'NaN', 'abc'])
def test_invalid_task_count(bad):
    with pytest.raises(ValueError): planned_count(getter({'uzk_sections': bad}), 'UZK')


def test_no_fabricated_weld():
    doc = Document()
    append_epb_protocol_weld_control(doc, 5, 'МК', lambda: None, [], 'годен')
    assert not doc.tables
    assert 'годен' not in '\n'.join(p.text for p in doc.paragraphs)


def test_scheme_shell_count_comes_from_passport():
    from scheme_ndt_overlays import extract_geometry
    geometry=extract_geometry({'vessel_elements':[{'name':f'Обечайка {i+1}', 'length_mm':2000 if i<5 else 1000} for i in range(6)]})
    assert geometry['shell_count']==6
    assert geometry['dimensions']['body_length_mm']==11000


def test_merged_header_preserves_column_widths():
    from epb_layout import format_table
    from docx.shared import Cm
    doc=Document();t=doc.add_table(rows=2,cols=3)
    t.cell(0,1).merge(t.cell(0,2))
    format_table(t,[3,2,4],header_rows=1)
    assert abs(t.cell(0,1).width-Cm(6))<1000
    assert abs(t.cell(1,1).width-Cm(2))<1000


def test_technical_uzk_does_not_include_surface_control():
    from form_template_filler import _fill_uzk_results
    doc=Document();t=doc.add_table(rows=2,cols=8)
    data={'weld_inspections':[{'control_method':'UZK','weld_number':'К1','conclusion':'годен'},
                              {'control_method':'MPK','weld_number':'К2','conclusion':'годен'}]}
    _fill_uzk_results(t,{'g':getter(data)})
    assert len(t.rows)==2
    assert 'К1' in ' '.join(c.text for c in t.rows[1].cells)
    assert 'К2' not in ' '.join(c.text for c in t.rows[1].cells)


def test_technical_hardness_supports_more_than_six_sections():
    from form_template_filler import _fill_hardness_matrix
    doc=Document();t=doc.add_table(rows=9,cols=7)
    tests=[{'weld_number':f'К{i+1}','hardness_weld':160+i} for i in range(8)]
    _fill_hardness_matrix(t,{'g':getter({'hardness_tests':tests})})
    assert len(t.rows)==11
    assert t.cell(10,1).text=='К8'
    assert t.cell(10,4).text=='167'


def test_contents_links_to_appendix_heading_not_table_reference():
    from form_template_filler import _refresh_vessel_contents
    from docx.oxml.ns import qn
    doc=Document();doc.add_paragraph('СОДЕРЖАНИЕ')
    doc.add_paragraph('ПРИЛОЖЕНИЕ № 1 Протокол9')
    doc.add_paragraph('1. Основания для проведения работ')
    doc.add_table(rows=1,cols=1).cell(0,0).text='ПРИЛОЖЕНИЕ № 1'
    actual=doc.add_paragraph('ПРИЛОЖЕНИЕ № 1 Протокол')
    _refresh_vessel_contents(doc)
    assert list(actual._p.iter(qn('w:bookmarkStart')))
    assert not list(doc.tables[0]._tbl.iter(qn('w:bookmarkStart')))


def test_source_snapshot_integrity(tmp_path):
    report = tmp_path / 'report.docx'; report.write_bytes(b'revision one')
    snapshot = tmp_path / 'snapshot.json'
    inputs = {'report_type': 'TECHNICAL_REPORT', 'inspection': {'data': {'working_pressure': 1}}, 'equipment': {}}
    snapshot.write_text(json.dumps({'snapshot_version': 1, 'inputs': inputs, 'report_sha256': hashlib.sha256(report.read_bytes()).hexdigest()}))
    assert load_technical_report_snapshot(report, snapshot) == inputs
    report.write_bytes(b'changed')
    with pytest.raises(ValueError, match='изменён'): load_technical_report_snapshot(report, snapshot)


def test_full_vessel_epb_has_reference_structure(tmp_path, monkeypatch):
    monkeypatch.setenv('DATABASE_URL', 'sqlite+aiosqlite:///:memory:')
    from word_generator import WordGenerator
    out = tmp_path / 'epb.docx'
    WordGenerator().generate_report_word(
        {'id': 'test', 'status': 'DRAFT', 'data': {'body_length_mm': 11000}},
        {'name': 'Сосуд', 'type_code': 'VESSEL', 'attributes': {}}, [], str(out), 'EXPERTISE')
    doc = Document(out)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
    assert any(h.startswith('8. Сведения') for h in headings)
    assert any(h.startswith('9. Вывод') for h in headings)
    assert [h.split()[1] for h in headings if h.startswith('Приложение ')] == list('АБВГД')
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert '72 точек' in text
    assert 'обеспечивает требуемый запас прочности' not in text
    # No uploaded PNG: the ordinary EPB path must render from source dimensions.
    from docx.oxml.ns import qn
    in_scheme = False
    drawings = 0
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith('Приложение В '):
            in_scheme = True
        elif paragraph.text.startswith('Приложение Г '):
            in_scheme = False
        if in_scheme:
            drawings += len(list(paragraph._p.iter(qn('w:drawing'))))
    assert drawings == 2  # VIK and UZT; no unrequested method maps.
    assert 'Схема контроля не предоставлена' not in text


@pytest.mark.parametrize('image', [None, 'missing-scheme.png'])
def test_epb_scheme_falls_back_to_constructor(tmp_path, image):
    from types import SimpleNamespace
    from vessel_epb_appendices import append_inspection_schemes
    data = {'control_scheme_image': image,
            'base_vessel_scheme': {'geometry': {'shell_count': 3}},
            'technical_task': {'uzk_sections': 2, 'pvk_sections': 2}}
    doc = Document()
    append_inspection_schemes(doc, SimpleNamespace(g=getter(data), inspection_data={'data': data}),
                             lambda *args, **kwargs: False)
    out = tmp_path / 'schemes.docx'
    doc.save(out)
    assert len(Document(out).inline_shapes) == 4
