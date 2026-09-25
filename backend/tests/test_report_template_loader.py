from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import pytest
from docx import Document
from report_template_loader import load_report_template


def test_legacy_template_conversion_is_isolated_and_preserves_source(tmp_path):
    source = tmp_path / "to-11.doc"
    source.write_bytes(b"legacy source")

    def convert(command, **kwargs):
        assert "docx:Office Open XML Text" in command
        assert any(arg.startswith("-env:UserInstallation=file:") for arg in command)
        output = Path(command[command.index("--outdir") + 1]) / "to-11.docx"
        doc = Document()
        doc.add_paragraph("Официальный шаблон")
        doc.save(output)
        return SimpleNamespace(returncode=0)

    with patch("report_template_loader._find_soffice", return_value="soffice"), patch(
        "report_template_loader.subprocess.run", side_effect=convert
    ):
        doc = load_report_template(source)
    assert doc.paragraphs[0].text == "Официальный шаблон"
    assert source.read_bytes() == b"legacy source"


def test_legacy_template_without_converter_fails_explicitly(tmp_path):
    with patch("report_template_loader._find_soffice", return_value=None):
        with pytest.raises(RuntimeError, match="LibreOffice"):
            load_report_template(tmp_path / "to-11.doc")
