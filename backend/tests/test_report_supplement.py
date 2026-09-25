from io import BytesIO
from docx import Document
from report_supplement import append_report_fields
import pytest
from report_forms_registry import FORMS_DIR


def test_additional_values_survive_word_save_with_zero_and_units():
    doc = Document()
    doc.add_paragraph("Основная форма")
    assert append_report_fields(doc, {"additional_data": {"report_fields": [
        {"section": "Электроизмерения", "label": "Сопротивление", "value": 0, "unit": "Ом"},
        {"label": "Замечание", "value": "Строка 1\nСтрока 2"},
        {}, None,
    ]}}) == 2
    output = BytesIO()
    doc.save(output)
    restored = Document(output)
    assert restored.paragraphs[0].text == "Основная форма"
    assert restored.tables[0].cell(1, 2).text == "0"
    assert restored.tables[0].cell(1, 3).text == "Ом"
    assert restored.tables[0].cell(2, 2).text == "Строка 1\nСтрока 2"


def test_empty_supplement_does_not_modify_official_form():
    doc = Document()
    before = doc.element.xml
    assert append_report_fields(doc, {"additional_data": {"report_fields": [{}]}}) == 0
    assert doc.element.xml == before


@pytest.mark.parametrize("template", sorted(FORMS_DIR.glob("to-*.docx")), ids=lambda p: p.stem)
def test_supplement_can_be_saved_in_each_official_docx_template(template):
    doc = Document(template)
    assert append_report_fields(doc, {"additional_data": {"report_fields": [
        {"label": "Контрольный показатель", "value": "123,45", "unit": "мм"},
    ]}}) == 1
    output = BytesIO()
    doc.save(output)
    assert Document(output).tables[-1].cell(1, 2).text == "123,45"
