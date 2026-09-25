"""Generate every registered form in an isolated directory; no database needed."""
import json
import sys
import tempfile
import argparse
from contextlib import nullcontext
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from report_forms_registry import load_forms_catalog, resolve_form_path
from report_template_loader import load_report_template
from form_template_filler_generic import fill_generic_official_form
from scripts.test_official_form_fillers import _base_data


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output-dir', type=Path)
    args = parser.parse_args()
    if args.output_dir:
        args.output_dir.mkdir(parents=True, exist_ok=True)
    results = []
    with (nullcontext(args.output_dir) if args.output_dir else tempfile.TemporaryDirectory(prefix="report_catalog_audit_")) as tmp:
        for meta in load_forms_catalog():
            fid = meta["id"]
            try:
                template = load_report_template(resolve_form_path(fid))
                target = Path(tmp) / f"{fid}.docx"
                fill_generic_official_form(
                    fid, {"date_performed": "2026-09-18", "data": _base_data()},
                    {"name": "Контрольный объект", "attributes": {}}, str(target),
                    org_settings={},
                )
                generated = Document(target)
                assert len(generated.tables) >= len(template.tables), "Template tables lost"
                results.append({"form": fid, "generated": True, "tables": len(generated.tables)})
            except Exception as exc:
                results.append({"form": fid, "generated": False, "error": str(exc)})
    print(json.dumps(results, ensure_ascii=False, indent=2))
    return int(any(not r["generated"] for r in results))


if __name__ == "__main__":
    raise SystemExit(main())
