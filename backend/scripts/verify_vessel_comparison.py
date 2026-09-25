"""Check generated examples against source snapshots and the reference structure."""
import json
import sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from epb_report_builder import EPB_TOC_ITEMS


def verify(root):
    results=[]
    for variant,length in [(1,11000),(2,5500)]:
        source=json.loads((root/f'fixture-{variant}.json').read_text(encoding='utf-8'))
        data=source['inspection']['data']
        epb=Document(root/f'{variant}_ЭПБ_{length}мм.docx')
        technical=Document(root/f'{variant}_Технический_отчет_{length}мм.docx')
        all_text=lambda d:' '.join(t.text or '' for t in d.element.iter(qn('w:t')))
        for d in [epb,technical]:
            text=all_text(d)
            assert data['vessel_name'] in text and data['serial_number'] in text
            assert '20102' not in text and '26-2335' not in text
            assert '09Г2С 09Г2С' not in text
        headings=[p.text for p in epb.paragraphs if p.style.name=='Heading 1']
        assert all(item in headings for item in EPB_TOC_ITEMS)
        uzt=next(t for t in epb.tables if len(t.columns)==12)
        actual={}
        for row in uzt.rows[2:]:
            cells=row.cells;name=cells[0].text
            for i in range(2,12,2):
                if cells[i].text:
                    key=(name,int(cells[i].text))
                    assert key not in actual
                    actual[key]=float(cells[i+1].text.replace(',','.'))
        expected={(p['location'],p['point_number']):p['thickness'] for p in data['thickness_measurements']}
        assert actual==expected
        shells=[e for e in data['vessel_elements'] if 'Обечайка' in e['name']]
        assert all(sum(k[0]==s['name'] for k in actual)==12 for s in shells)
        uzk=next(t for t in epb.tables if len(t.columns)==8 and t.cell(0,0).text=='Номер соединения')
        hardness=next(t for t in epb.tables if len(t.columns)==8 and t.cell(0,0).text=='Соединение / участок')
        surface=next(t for t in epb.tables if t.cell(0,0).text=='Номер соединения / участка')
        task=data['technical_task']
        assert len(uzk.rows)-1==task['uzk_sections']
        assert len(hardness.rows)-3==task['hardness_sections']
        assert len(surface.rows)-1==task['mpk_sections']+task['pvk_sections']
        assert len(technical.tables[31].rows)-1==task['uzk_sections']
        assert len(technical.tables[36].rows)-1==task['mpk_sections']+task['pvk_sections']
        results.append(dict(variant=variant,body_length_mm=length,shells=len(shells),
                            shell_points=12*len(shells),all_measurements=len(actual),
                            task=task,all_source_measurements_match=True,
                            sections_1_to_9_and_appendices_A_to_D=True))
    return results


if __name__=='__main__':
    root=Path(sys.argv[1]);result=verify(root)
    (root/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2))
